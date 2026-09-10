from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

import config
from data.model import SitRepData
from docx import Document
from docx.document import Document as DocumentT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reporting.highlights import (
    build_actions_laboratoire,
    build_actions_surveillance,
    build_commentaire_tableau1,
    build_commentaire_tableau2,
    build_conclusion,
    build_faits_saillants,
    build_resume_points_cles,
)
from utils import docx as dx
from utils.dates import fr_date as _fr_date
from utils.numbers import fmt_int as _fr_int

Filler = Callable[[dx.Cursor], None]

_para = dx.para
_bullet = dx.bullet
_table = dx.table
_province_zone_table = dx.province_zone_table
_replace_marker = dx.replace_marker
_set_inline_marker = dx.set_inline_marker
_fill_shape_lines = dx.fill_shape_lines
_normalize_markers = dx.normalize_markers


def _fr_pct(value: float | str, decimals: int = 1) -> str:
    """Formate un pourcentage à la française (virgule décimale).

    Returns:
        str: ``"25,3%"`` ou la valeur telle quelle si non numérique (``ND``).
    """
    if not isinstance(value, (int, float)):
        return str(value)
    return f"{value:.{decimals}f}".replace(".", ",") + "%"


def _pct_of(k: object, total: object) -> str:
    if not isinstance(k, int) or not isinstance(total, int) or not total:
        return config.ND
    return _fr_pct(k / total * 100)


def _zones_fraction(touchees: int, total: int | None) -> str:
    """« 28/36 (77,8 %) », ou juste ``touchees`` si le total est inconnu."""
    if not total:
        return str(touchees)
    return f"{touchees}/{total} ({_pct_of(touchees, total)})"


def _set_document_language(doc: DocumentT, lang: str = "fr-FR") -> None:
    """Définit le français comme langue d'édition/correction par défaut.

    Reprise de v1 (``sitrep/code/generate_sitrep/reporting/render.py``) —
    aucune dépendance au schéma de données.
    """
    styles = doc.styles.element
    node = styles.find(qn("w:docDefaults"))
    if node is None:
        node = OxmlElement("w:docDefaults")
        styles.insert(0, node)
    for tag in ("w:rPrDefault", "w:rPr", "w:lang"):
        child = node.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            node.append(child)
        node = child
    node.set(qn("w:val"), lang)
    node.set(qn("w:eastAsia"), lang)


def _image_filler(doc: DocumentT, path: str | Path | None, *, width_in: float = 6.0) -> Filler:
    """Insère une image centrée sous le marqueur (corps de document).

    Écrit « (visuel indisponible) » si le fichier est absent, pour ne pas
    casser le rendu.

    Returns:
        Filler: La fonction de remplissage à passer à ``replace_marker``.
    """

    def fill(cur: dx.Cursor) -> None:
        if not path or not Path(path).exists():
            cur.add(_para(doc, "(visuel indisponible)", italic=True, size=9, color="BFBFBF", align=dx.CENTER)._p)
            return
        p = _para(doc, "", align=dx.CENTER)
        p.add_run().add_picture(str(path), width=dx.Inches(width_in))
        cur.add(p._p)

    return fill


# --- Marqueurs mono-valeur (corps ou forme) ---------------------------------


def _fill_title_identity(doc: DocumentT, data: SitRepData) -> None:
    """``[[TITRE_NUMERO]]``/``[[DATE_RAPPORTAGE]]``/``[[DATE_PUBLICATION]]`` (corps).

    Contrairement à v1 (``_fill_title_number``, où le marqueur portait le
    texte complet « SitRep N°… »), le nouveau template a « SitRep N° » en
    texte statique dans un run précédent : le marqueur ne doit porter que le
    numéro/la date, sous peine de doublon (vérifié par test isolé).

    ``[[DATE_RAPPORTAGE]]`` affiche toujours une seule date (``reporting_end``),
    pas la plage ``reporting_label`` (utilisée ailleurs pour le libellé de
    période « jour »/« 24h »).
    """
    titre = f"{data.sitrep_number}/{config.INCIDENT}_{data.reporting_end:%d/%m/%Y}"
    if data.scope_label:
        titre = f"{titre} — {data.scope_label}"
    _set_inline_marker(doc, "[[TITRE_NUMERO]]", titre)
    _set_inline_marker(doc, "[[DATE_RAPPORTAGE]]", _fr_date(data.reporting_end))
    _set_inline_marker(doc, "[[DATE_PUBLICATION]]", _fr_date(data.publication_date))


def _fill_kpi(doc: DocumentT, data: SitRepData) -> None:
    """Les 4 cartes KPI (forme) : ``[[KPI_CUMUL_*]]``/``[[KPI_LETALITE]]``."""
    kpi = data.kpi
    _set_inline_marker(doc, "[[KPI_CUMUL_CONFIRMES]]", _fr_int(kpi["cumul_confirmes"]))
    _set_inline_marker(doc, "[[KPI_CUMUL_DECES]]", _fr_int(kpi["cumul_deces"]))
    _set_inline_marker(doc, "[[KPI_CUMUL_GUERIS]]", _fr_int(kpi["cumul_gueris"]))
    _set_inline_marker(doc, "[[KPI_LETALITE]]", _fr_pct(kpi["letalite"]))


def _fill_provinces_zs_as(doc: DocumentT, data: SitRepData) -> None:
    """Provinces/ZS/AS touchées (forme)."""
    n_provinces = len(data.provinces_touchees)
    n_zs = sum(int(d.get("touchees", 0)) for d in data.zones_atteintes.values())
    aires = data.aires_atteintes

    _set_inline_marker(doc, "[[NB_PROVINCES_TOUCHEES]]", _fr_int(n_provinces))
    _set_inline_marker(doc, "[[LISTE_PROVINCES_TOUCHEES]]", ", ".join(data.provinces_touchees))
    _set_inline_marker(doc, "[[NB_ZS_TOUCHEES]]", _fr_int(n_zs))
    _set_inline_marker(doc, "[[TOTAL_ZS]]", _fr_int(config.TOTAL_ZONES_SANTE))
    _set_inline_marker(doc, "[[PCT_ZS_TOUCHEES]]", _pct_of(n_zs, config.TOTAL_ZONES_SANTE).rstrip("%"))
    _set_inline_marker(doc, "[[NB_AS_TOUCHEES]]", _fr_int(aires.get("touchees", config.ND)))
    _set_inline_marker(doc, "[[TOTAL_AS]]", _fr_int(aires.get("total", config.ND)))
    _set_inline_marker(
        doc, "[[PCT_AS_TOUCHEES]]", _pct_of(aires.get("touchees"), aires.get("total")).rstrip("%")
    )


# --- Marqueurs multi-lignes en forme -----------------------------------


def _fill_faits_saillants(doc: DocumentT, data: SitRepData) -> None:
    _fill_shape_lines(doc, "[[FAITS_SAILLANTS]]", build_faits_saillants(data))


def _fill_resume_points_cles(doc: DocumentT, data: SitRepData) -> None:
    _fill_shape_lines(doc, "[[RESUME_POINTS_CLES]]", build_resume_points_cles(data))


def _fill_conclusion(doc: DocumentT, data: SitRepData, narrative: dict) -> None:
    recommandation = narrative.get("conclusion_recommandation")
    _fill_shape_lines(doc, "[[CONCLUSION]]", build_conclusion(data, recommandation))


# --- Figures (corps) : [[COURBE_EPI]] / [[CARTE_1]] / [[CARTE_2]] / [[PYRAMIDE]] --


def _fill_figures(doc: DocumentT, charts: dict[str, Path | None]) -> None:
    _replace_marker(doc, "[[COURBE_EPI]]", _image_filler(doc, charts.get("epi_curve"), width_in=6.5))
    _replace_marker(doc, "[[CARTE_1]]", _image_filler(doc, charts.get("zone_situation_map_cumul"), width_in=6.5))
    _replace_marker(doc, "[[CARTE_2]]", _image_filler(doc, charts.get("zone_situation_map_jour"), width_in=6.5))
    _replace_marker(doc, "[[PYRAMIDE]]", _image_filler(doc, charts.get("age_sex_pyramid"), width_in=6.5))


# --- Tableaux + commentaires (corps) : replace_marker classique -------------


def _fill_tableaux(doc: DocumentT, data: SitRepData) -> None:
    def fill_tableau1(cur: dx.Cursor) -> None:
        rows = [
            [
                r["province"],
                _fr_int(r["nouveaux"]),
                _fr_int(r["confirmes"]),
                _fr_int(r["deces"]),
                _fr_pct(r["cfr"]),
                _zones_fraction(r["zones_touchees"], r["zones_total"]),
            ]
            for r in data.tableau1
        ]
        t = data.tableau1_total
        rows.append(
            [
                "Total",
                _fr_int(t["nouveaux"]),
                _fr_int(t["confirmes"]),
                _fr_int(t["deces"]),
                _fr_pct(t["cfr"]),
                _zones_fraction(t["zones_touchees"], t["zones_total"]),
            ]
        )
        cur.add(
            _table(
                doc,
                [
                    "Province",
                    "Nouveaux cas confirmés (24h)",
                    "Cas confirmés",
                    "Décès (confirmés)",
                    "Létalité",
                    "Zones de santé touchées",
                ],
                rows,
                table_title=(
                    "Tableau 1. Répartition des cas et décès confirmés par province touchée, "
                    f"au {_fr_date(data.reporting_end)}."
                ),
                table_title_fill=config.ACCENT_RED,
                table_title_font_color="FFFFFF",
                total_fill=config.ACCENT_RED,
                total_font_color="FFFFFF",
            )._tbl
        )

    def fill_tableau2(cur: dx.Cursor) -> None:
        rows = [
            {
                "label": r["province"] if r["is_province"] else r["zone"],
                "is_province": r["is_province"],
                "values": [
                    _fr_int(r["confirmes"]),
                    _fr_int(r["deces"]),
                    _fr_pct(r["cfr"]),
                    _fr_int(r["nouveaux"]),
                    _fr_int(r["deces_communautaires"]),
                    _fr_int(r["deces_intra_cte"]),
                    _fr_int(r["deces_total_jour"]),
                ],
            }
            for r in data.tableau2
        ]
        t = data.tableau2_total
        rows.append(
            {
                "label": "Total",
                "is_province": False,
                "values": [
                    _fr_int(t["confirmes"]),
                    _fr_int(t["deces"]),
                    _fr_pct(t["cfr"]),
                    _fr_int(t["nouveaux"]),
                    _fr_int(t["deces_communautaires"]),
                    _fr_int(t["deces_intra_cte"]),
                    _fr_int(t["deces_total_jour"]),
                ],
            }
        )
        cur.add(
            _province_zone_table(
                doc,
                [("Nombre cumulatif", 3), ("Situation du jour (24h)", 4)],
                [
                    "Cas (n)",
                    "Décès (n)",
                    "Létalité",
                    "Nouv. Cas",
                    "Décès communautaires (Swab+)",
                    "Décès confirmés intra-CTE",
                    "Total décès confirmés",
                ],
                rows,
                table_title=(
                    "Tableau 2. Répartition des cas et décès confirmés par province et zone de "
                    f"santé, au {_fr_date(data.reporting_end)}."
                ),
                header_fill=config.ACCENT_DARK_BLUE,
                header_font_color="FFFFFF",
            )._tbl
        )

    def fill_commentaire1(cur: dx.Cursor) -> None:
        for line in build_commentaire_tableau1(data):
            cur.add(_bullet(doc, line)._p)

    def fill_commentaire2(cur: dx.Cursor) -> None:
        for line in build_commentaire_tableau2(data):
            cur.add(_bullet(doc, line)._p)

    _replace_marker(doc, "[[TABLEAU_1]]", fill_tableau1)
    _replace_marker(doc, "[[COMMENTAIRE_TABLEAU_1]]", fill_commentaire1)
    _replace_marker(doc, "[[TABLEAU_2]]", fill_tableau2)
    _replace_marker(doc, "[[COMMENTAIRE_TABLEAU_2]]", fill_commentaire2)


def _inject_narrative(doc: DocumentT, data: SitRepData, narrative: dict) -> None:
    """Marqueurs narratifs : calculés (surveillance/labo, corps) ou manuels (contacts, forme).

    Les 3 marqueurs contacts vivent désormais à l'intérieur d'une forme dans
    le nouveau template (v2) — plus en corps de document comme dans
    l'ancien (cf. script lxml : ``[[CONTACTS_MANAGER]]`` etc. sont dans un
    ``w:txbxContent``) — d'où ``fill_shape_lines`` (cf. B.6) au lieu de
    ``replace_marker``, qui ne cible que les paragraphes de corps.
    """
    narrative = narrative or {}
    contacts = narrative.get("contacts", {}) or {}

    def bullets(items: list[str] | None) -> Filler:
        def fill(cur: dx.Cursor) -> None:
            if not items:
                cur.add(_para(doc, "À compléter.", italic=True, size=10, color="808080")._p)
                return
            for it in items:
                cur.add(_bullet(doc, str(it))._p)

        return fill

    _replace_marker(doc, "[[ACTIONS_SURVEILLANCE]]", bullets(build_actions_surveillance(data)))
    _replace_marker(doc, "[[ACTIONS_LABORATOIRE]]", bullets(build_actions_laboratoire(data)))
    _fill_shape_lines(doc, "[[CONTACTS_MANAGER]]", contacts.get("manager"))
    _fill_shape_lines(doc, "[[CONTACTS_OPERATIONS]]", contacts.get("operations"))
    _fill_shape_lines(doc, "[[CONTACTS_REDACTIONS]]", contacts.get("redactions"))


def render(
    data: SitRepData,
    charts: dict[str, Path | None],
    template_path: str | Path,
    output_path: str | Path,
    narrative: dict | None = None,
) -> Path:
    """Produit le fichier SitRep .docx et renvoie son chemin.

    Returns:
        Path: Le chemin du ``.docx`` généré.
    """
    doc = Document(str(template_path))
    narrative = narrative or {}

    _normalize_markers(doc)
    _set_document_language(doc)
    _fill_title_identity(doc, data)
    _fill_kpi(doc, data)
    _fill_provinces_zs_as(doc, data)
    _fill_faits_saillants(doc, data)
    _fill_resume_points_cles(doc, data)
    _fill_conclusion(doc, data, narrative)
    _fill_figures(doc, charts)
    _fill_tableaux(doc, data)
    _inject_narrative(doc, data, narrative)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
